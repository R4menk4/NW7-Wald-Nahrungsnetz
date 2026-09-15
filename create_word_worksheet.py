from pathlib import Path
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path('output/docx');out.mkdir(parents=True,exist_ok=True)
tmp=Path('tmp/docx');tmp.mkdir(parents=True,exist_ok=True)
pdf=pdfium.PdfDocument('dist/assets/arbeitsblatt-baum.pdf')
for i,box,name in [(0,(22,64,138,141.33),'baum'),(1,(19,65,126,136.33),'wurzel')]:
 im=pdf[i].render(scale=3).to_pil()
 k=3*72/25.4
 im.crop(tuple(round(v*k) for v in box)).save(tmp/(name+'.png'))
d=Document();s=d.sections[0]
s.page_width=Mm(210);s.page_height=Mm(297)
s.top_margin=Mm(14);s.bottom_margin=Mm(15);s.left_margin=Mm(17);s.right_margin=Mm(17)
s.footer_distance=Mm(7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
st=d.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(20);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(7)
d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].paragraph_format.space_before=Pt(10);d.styles['Heading 1'].paragraph_format.space_after=Pt(4)
def p(text='',size=None,after=5,bold=False):
 a=d.add_paragraph();a.paragraph_format.space_after=Pt(after)
 r=a.add_run(text);r.bold=bold
 if size:r.font.size=Pt(size)
 return a
def header():
 p('ÖKOSYSTEM WALD',10,after=4,bold=True)
 d.add_paragraph('Wurzel Stamm und Krone','Title')
 p('Name: __________________________________    Datum: __________________',10,after=7)
def picture(name,width):
 a=d.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_after=Pt(3)
 shape=a.add_run().add_picture(str(tmp/(name+'.png')),width=Mm(width))
 shape._inline.docPr.set('descr','Baum mit Markierungen A Krone B Stamm C Wurzeln' if name=='baum' else 'Wurzelspitze mit Markierungen A Wurzelhaare B Wurzelhaube')
def table_borders(t):
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');b.append(z)
 pr.append(b)
 for row in t.rows:
  for cell in row.cells:
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   mar=OxmlElement('w:tcMar')
   for edge,v in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');mar.append(e)
   cell._tc.get_or_add_tcPr().append(mar)
header()
d.add_paragraph('1  Den Baum erkunden','Heading 1')
p('Öffne im Lernprogramm die drei Bereiche. Beschrifte A, B und C.')
picture('baum',111)
p('A: ____________________   B: ____________________   C: ____________________',10,after=5)
d.add_paragraph('2  Den Grundbauplan ergänzen','Heading 1')
p('Der Baum besteht aus Wurzeln und dem ______________________________.',after=9)
p('Zum Spross gehören der __________________ und die __________________.',after=9)
p('Die Krone besteht aus ________________, ________________ und Blättern.',after=5)
d.add_paragraph('3  Die Aufgaben festhalten','Heading 1')
p('Notiere nach der digitalen Zuordnung jeweils zwei Aufgaben.')
t=d.add_table(rows=4,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
t.columns[0].width=Mm(43);t.columns[1].width=Mm(133)
for row in t.rows:
 row.cells[0].width=Mm(43);row.cells[1].width=Mm(133)
for i,name in enumerate(['Teil des Baumes','Wurzeln','Stamm','Laubblätter']):
 t.cell(i,0).text=name;t.cell(i,1).text='Zwei Aufgaben' if i==0 else '____________________________________________________\n____________________________________________________'
 for j in range(2):
  for para in t.cell(i,j).paragraphs:
   para.paragraph_format.space_after=Pt(2)
   for r in para.runs:r.font.size=Pt(10);r.bold=(i==0 or j==0)
 if i:
  t.rows[i].height=Mm(16);t.rows[i].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
 else:
  for cell in t.rows[0].cells:
   e=OxmlElement('w:shd');e.set(qn('w:fill'),'E7EBED');cell._tc.get_or_add_tcPr().append(e)
  t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
table_borders(t)
d.add_page_break()
header()
d.add_paragraph('4  Die Wurzel unter der Lupe','Heading 1')
p('Bearbeite die Wurzellupe im Lernprogramm. Beschrifte A und B.')
picture('wurzel',106)
p('A: ______________________________    B: ______________________________',after=9)
p('Viele Wurzelhaare vergrößern die ________________________ der Wurzel.',after=9)
p('Dadurch kann sie Wasser und gelöste ______________________ aufnehmen.',after=9)
p('Die Wurzelhaube schützt die _________________________________________.',after=5)
d.add_paragraph('5  Richtig oder falsch','Heading 1')
p('Kreuze an. Prüfe digital und berichtige die beiden falschen Aussagen darunter.',10,after=7)
statements=['Die Wurzeln halten den Baum im Boden.','Die Borke nimmt Wasser aus dem Boden auf.','Der Stamm trägt die Krone und transportiert Stoffe.','Zur Krone gehören nur die Blätter.','Die Wurzelhaube schützt die empfindliche Wurzelspitze.']
t=d.add_table(rows=6,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
for col,w in zip(t.columns,[140,18,18]):col.width=Mm(w)
for row in t.rows:
 for cell,w in zip(row.cells,[140,18,18]):cell.width=Mm(w)
for j,x in enumerate(['Aussage','richtig','falsch']):t.cell(0,j).text=x
for i,x in enumerate(statements,1):
 t.cell(i,0).text=str(i)+'. '+x
 for j in (1,2):t.cell(i,j).text='☐'
for i,row in enumerate(t.rows):
 for j,cell in enumerate(row.cells):
  for a in cell.paragraphs:
   a.paragraph_format.space_after=Pt(2)
   if j:a.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for r in a.runs:
    r.font.size=Pt(9.5 if j==0 else 10)
    if i==0:r.bold=True
    elif j:r.font.name='Segoe UI Symbol'
for cell in t.rows[0].cells:
 e=OxmlElement('w:shd');e.set(qn('w:fill'),'E7EBED');cell._tc.get_or_add_tcPr().append(e)
table_borders(t)
p('So lauten die falschen Aussagen richtig:',10,after=6,bold=True)
for _ in range(3):p('________________________________________________________________________',10,after=6)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Ökosystem Wald  |  ');r.font.name='Arial';r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.title='Wurzel Stamm und Krone'
d.core_properties.author='Wald entdecken'
path=out/'Arbeitsblatt_Wurzel_Stamm_und_Krone.docx'
d.save(path)
print(path.resolve())

